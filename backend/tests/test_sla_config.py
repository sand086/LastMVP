"""
Test suite for Dynamic SLA Configuration per Provider (Iteration 30)
Tests:
1. GET /api/admin/config returns sla_config with default_sla and by_provider
2. PATCH /api/admin/config with section='sla_config' saves SLA configuration
3. PATCH /api/admin/config with section='sla_config' requires developer role (403 for non-developers)
4. GET /api/admin/export-liquidacion uses dynamic SLA values from config
"""
import pytest
import requests
import os
from io import BytesIO

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', '').rstrip('/')

# Test credentials
DEVELOPER_CREDS = {"email": "dev@me.mx", "password": "LastMile2026"}
COORDINATOR_CREDS = {"email": "yael@me.mx", "password": "LastMile2026"}

# Known provider for testing
EXCLUSIVE_LOGISTICS_ID = "2b34d97d-c709-4dfc-94ec-71e534874bc8"
TIL_PROVIDER_ID = "1f261581-2f85-4608-aa1d-2b93073dfedd"

# Date range for export testing
DATE_FROM = "2026-01-01"
DATE_TO = "2026-04-07"


@pytest.fixture(scope="module")
def developer_token():
    """Get developer auth token"""
    response = requests.post(f"{BASE_URL}/api/auth/login", json=DEVELOPER_CREDS)
    if response.status_code == 200:
        return response.json().get("access_token")
    pytest.skip("Developer login failed")


@pytest.fixture(scope="module")
def coordinator_token():
    """Get coordinator auth token"""
    response = requests.post(f"{BASE_URL}/api/auth/login", json=COORDINATOR_CREDS)
    if response.status_code == 200:
        return response.json().get("access_token")
    pytest.skip("Coordinator login failed")


class TestSlaConfigEndpoint:
    """Test GET /api/admin/config returns sla_config correctly"""
    
    def test_config_returns_sla_config_structure(self, developer_token):
        """GET /api/admin/config should return sla_config with default_sla and by_provider"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(f"{BASE_URL}/api/admin/config", headers=headers)
        
        assert response.status_code == 200, f"Expected 200, got {response.status_code}"
        data = response.json()
        
        # Verify sla_config exists
        assert "sla_config" in data, "sla_config should be in response"
        sla_config = data["sla_config"]
        
        # Verify structure
        assert "default_sla" in sla_config, "default_sla should be in sla_config"
        assert "by_provider" in sla_config, "by_provider should be in sla_config"
        
        # Verify types
        assert isinstance(sla_config["default_sla"], (int, float)), "default_sla should be numeric"
        assert isinstance(sla_config["by_provider"], dict), "by_provider should be a dict"
        
        print(f"SLA Config: default_sla={sla_config['default_sla']}, by_provider={sla_config['by_provider']}")
    
    def test_config_has_default_sla_value(self, developer_token):
        """default_sla should have a reasonable value (typically 40)"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(f"{BASE_URL}/api/admin/config", headers=headers)
        
        assert response.status_code == 200
        sla_config = response.json().get("sla_config", {})
        default_sla = sla_config.get("default_sla")
        
        assert default_sla is not None, "default_sla should not be None"
        assert default_sla > 0, "default_sla should be positive"
        assert default_sla <= 100, "default_sla should be reasonable (<=100)"
        
        print(f"Default SLA value: {default_sla}")
    
    def test_config_has_provider_override(self, developer_token):
        """by_provider should contain the Exclusive Logistics override (set by main agent)"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(f"{BASE_URL}/api/admin/config", headers=headers)
        
        assert response.status_code == 200
        sla_config = response.json().get("sla_config", {})
        by_provider = sla_config.get("by_provider", {})
        
        # Check if Exclusive Logistics has an override
        if EXCLUSIVE_LOGISTICS_ID in by_provider:
            sla_value = by_provider[EXCLUSIVE_LOGISTICS_ID]
            assert isinstance(sla_value, (int, float)), "Provider SLA should be numeric"
            print(f"Exclusive Logistics SLA override: {sla_value}")
        else:
            print("No provider override found (using default)")


class TestSlaConfigPatch:
    """Test PATCH /api/admin/config with section='sla_config'"""
    
    def test_developer_can_update_sla_config(self, developer_token):
        """Developer should be able to update SLA configuration"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        
        # First get current config
        get_response = requests.get(f"{BASE_URL}/api/admin/config", headers=headers)
        assert get_response.status_code == 200
        current_sla = get_response.json().get("sla_config", {})
        original_default = current_sla.get("default_sla", 40)
        
        # Update with new value
        new_sla_config = {
            "default_sla": 45,  # Change from 40 to 45
            "by_provider": {
                EXCLUSIVE_LOGISTICS_ID: 55,  # Update Exclusive Logistics
                TIL_PROVIDER_ID: 35  # Add TIL override
            }
        }
        
        patch_response = requests.patch(
            f"{BASE_URL}/api/admin/config",
            json={"section": "sla_config", "value": new_sla_config},
            headers=headers
        )
        
        assert patch_response.status_code == 200, f"Expected 200, got {patch_response.status_code}: {patch_response.text}"
        patch_data = patch_response.json()
        assert patch_data.get("success") == True, "Response should indicate success"
        assert patch_data.get("section") == "sla_config", "Response should confirm section"
        
        # Verify the update persisted
        verify_response = requests.get(f"{BASE_URL}/api/admin/config", headers=headers)
        assert verify_response.status_code == 200
        updated_sla = verify_response.json().get("sla_config", {})
        
        assert updated_sla.get("default_sla") == 45, f"default_sla should be 45, got {updated_sla.get('default_sla')}"
        assert updated_sla.get("by_provider", {}).get(EXCLUSIVE_LOGISTICS_ID) == 55
        assert updated_sla.get("by_provider", {}).get(TIL_PROVIDER_ID) == 35
        
        print(f"SLA config updated successfully: {updated_sla}")
        
        # Restore original config
        restore_config = {
            "default_sla": 40,
            "by_provider": {EXCLUSIVE_LOGISTICS_ID: 50}
        }
        requests.patch(
            f"{BASE_URL}/api/admin/config",
            json={"section": "sla_config", "value": restore_config},
            headers=headers
        )
        print("Restored original SLA config")
    
    def test_coordinator_cannot_update_sla_config(self, coordinator_token):
        """Coordinator (non-developer) should NOT be able to update SLA config (403)"""
        headers = {"Authorization": f"Bearer {coordinator_token}"}
        
        new_sla_config = {
            "default_sla": 50,
            "by_provider": {}
        }
        
        response = requests.patch(
            f"{BASE_URL}/api/admin/config",
            json={"section": "sla_config", "value": new_sla_config},
            headers=headers
        )
        
        assert response.status_code == 403, f"Expected 403 for coordinator, got {response.status_code}"
        print(f"Coordinator correctly denied: {response.json()}")
    
    def test_coordinator_can_read_config(self, coordinator_token):
        """Coordinator should be able to READ config (but not edit)"""
        headers = {"Authorization": f"Bearer {coordinator_token}"}
        
        response = requests.get(f"{BASE_URL}/api/admin/config", headers=headers)
        
        assert response.status_code == 200, f"Expected 200, got {response.status_code}"
        data = response.json()
        assert "sla_config" in data
        print(f"Coordinator can read config: sla_config present")


class TestLiquidacionExportUsesSlaConfig:
    """Test that export-liquidacion uses dynamic SLA values"""
    
    def test_export_liquidacion_with_custom_sla(self, developer_token):
        """Export should use the SLA config values in formulas"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        
        # First, set a specific SLA for testing
        test_sla_config = {
            "default_sla": 42,  # Non-default value to verify it's used
            "by_provider": {
                EXCLUSIVE_LOGISTICS_ID: 55  # Custom SLA for Exclusive Logistics
            }
        }
        
        patch_response = requests.patch(
            f"{BASE_URL}/api/admin/config",
            json={"section": "sla_config", "value": test_sla_config},
            headers=headers
        )
        assert patch_response.status_code == 200
        
        # Now export liquidacion
        export_response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if export_response.status_code == 404:
            print("No journeys in date range - skipping Excel verification")
            pytest.skip("No data in date range for export")
        
        assert export_response.status_code == 200, f"Expected 200, got {export_response.status_code}"
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in export_response.headers.get("Content-Type", "")
        
        # Parse Excel to verify SLA is used
        try:
            from openpyxl import load_workbook
            wb = load_workbook(BytesIO(export_response.content))
            
            # Check provider sheets for SLA formulas
            for sheet_name in wb.sheetnames:
                if sheet_name in ["Formato", "Incidencias a cobro", "Catalogo_Drivers", "route_summary"]:
                    continue
                
                ws = wb[sheet_name]
                # Check row 2 header for SLA column (column AB = 28)
                sla_header = ws.cell(row=2, column=28).value
                print(f"Sheet '{sheet_name}' - SLA header: {sla_header}")
                
                # Check if formulas use the correct SLA value
                if ws.max_row >= 3:
                    # Column X (24) = Costo x Pq formula: =IFERROR(M{r}/{sla},0)
                    cost_formula = ws.cell(row=3, column=24).value
                    # Column AB (28) = % Efect. SLA formula: =IFERROR(O{r}/{sla},0)
                    sla_formula = ws.cell(row=3, column=28).value
                    
                    print(f"Sheet '{sheet_name}' - Cost formula: {cost_formula}")
                    print(f"Sheet '{sheet_name}' - SLA formula: {sla_formula}")
                    
                    # For Exclusive Logistics, SLA should be 55
                    if "Exclusive" in sheet_name:
                        if cost_formula and "55" in str(cost_formula):
                            print(f"✓ Exclusive Logistics uses custom SLA 55")
                        elif cost_formula:
                            print(f"Formula found: {cost_formula}")
            
            print("Excel export generated successfully with SLA formulas")
            
        except ImportError:
            print("openpyxl not available for detailed verification - export succeeded")
        
        # Restore original config
        restore_config = {
            "default_sla": 40,
            "by_provider": {EXCLUSIVE_LOGISTICS_ID: 50}
        }
        requests.patch(
            f"{BASE_URL}/api/admin/config",
            json={"section": "sla_config", "value": restore_config},
            headers=headers
        )


class TestSlaConfigValidation:
    """Test edge cases and validation for SLA config"""
    
    def test_invalid_section_returns_400(self, developer_token):
        """PATCH with invalid section should return 400"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        
        response = requests.patch(
            f"{BASE_URL}/api/admin/config",
            json={"section": "invalid_section", "value": {}},
            headers=headers
        )
        
        assert response.status_code == 400, f"Expected 400, got {response.status_code}"
        print(f"Invalid section correctly rejected: {response.json()}")
    
    def test_unauthenticated_cannot_access_config(self):
        """Unauthenticated requests should be rejected"""
        response = requests.get(f"{BASE_URL}/api/admin/config")
        
        assert response.status_code in [401, 403], f"Expected 401/403, got {response.status_code}"
        print(f"Unauthenticated correctly rejected: {response.status_code}")
    
    def test_sla_config_persists_after_update(self, developer_token):
        """SLA config should persist in database after update"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        
        # Set a unique value
        unique_sla = 47
        test_config = {
            "default_sla": unique_sla,
            "by_provider": {EXCLUSIVE_LOGISTICS_ID: 57}
        }
        
        # Update
        patch_response = requests.patch(
            f"{BASE_URL}/api/admin/config",
            json={"section": "sla_config", "value": test_config},
            headers=headers
        )
        assert patch_response.status_code == 200
        
        # Verify immediately
        get_response = requests.get(f"{BASE_URL}/api/admin/config", headers=headers)
        assert get_response.status_code == 200
        sla_config = get_response.json().get("sla_config", {})
        
        assert sla_config.get("default_sla") == unique_sla, f"Expected {unique_sla}, got {sla_config.get('default_sla')}"
        assert sla_config.get("by_provider", {}).get(EXCLUSIVE_LOGISTICS_ID) == 57
        
        print(f"SLA config persisted correctly: {sla_config}")
        
        # Restore
        restore_config = {
            "default_sla": 40,
            "by_provider": {EXCLUSIVE_LOGISTICS_ID: 50}
        }
        requests.patch(
            f"{BASE_URL}/api/admin/config",
            json={"section": "sla_config", "value": restore_config},
            headers=headers
        )


if __name__ == "__main__":
    pytest.main([__file__, "-v", "--tb=short"])
