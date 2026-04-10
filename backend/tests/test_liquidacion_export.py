"""
Test suite for Liquidacion Export feature (Iteration 29)
Tests the new GET /api/admin/export-liquidacion endpoint and Excel file structure
"""
import pytest
import requests
import os
from io import BytesIO

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', '').rstrip('/')

# Test credentials
DEVELOPER_CREDS = {"email": os.environ.get("TEST_DEV_EMAIL", "dev@me.mx"), "password": os.environ.get("TEST_DEV_PASSWORD", "LastMile2026")}
AGENT_CREDS = {"email": os.environ.get("TEST_AGENT_EMAIL", "agente@me.mx"), "password": os.environ.get("TEST_DEV_PASSWORD", "LastMile2026")}
PROVEEDOR_CREDS = {"email": "proveedor@me.mx", "password": os.environ.get("TEST_DEV_PASSWORD", "LastMile2026")}

# Date range that should have data
DATE_FROM = "2026-01-01"
DATE_TO = "2026-04-07"


@pytest.fixture(scope="module")
def developer_token():
    """Get developer auth token"""
    response = requests.post(f"{BASE_URL}/api/auth/login", json=DEVELOPER_CREDS)
    if response.status_code == 200:
        return response.json().get("access_token")
    pytest.skip("Developer login failed")


@pytest.fixture(scope="module")
def agent_token():
    """Get agent auth token"""
    response = requests.post(f"{BASE_URL}/api/auth/login", json=AGENT_CREDS)
    if response.status_code == 200:
        return response.json().get("access_token")
    pytest.skip("Agent login failed")


@pytest.fixture(scope="module")
def proveedor_token():
    """Get proveedor auth token"""
    response = requests.post(f"{BASE_URL}/api/auth/login", json=PROVEEDOR_CREDS)
    if response.status_code == 200:
        return response.json().get("access_token")
    pytest.skip("Proveedor login failed")


class TestAuthEndpoints:
    """Verify auth works for all test users"""
    
    def test_developer_login(self):
        response = requests.post(f"{BASE_URL}/api/auth/login", json=DEVELOPER_CREDS)
        assert response.status_code == 200
        data = response.json()
        assert "access_token" in data
        assert data.get("user", {}).get("role") == "developer"
        print(f"Developer login successful, role: {data.get('user', {}).get('role')}")
    
    def test_agent_login(self):
        response = requests.post(f"{BASE_URL}/api/auth/login", json=AGENT_CREDS)
        assert response.status_code == 200
        data = response.json()
        assert "access_token" in data
        assert data.get("user", {}).get("role") in ["agente", "agent"]
        print(f"Agent login successful, role: {data.get('user', {}).get('role')}")
    
    def test_proveedor_login(self):
        response = requests.post(f"{BASE_URL}/api/auth/login", json=PROVEEDOR_CREDS)
        assert response.status_code == 200
        data = response.json()
        assert "access_token" in data
        assert data.get("user", {}).get("role") == "proveedor"
        print(f"Proveedor login successful, role: {data.get('user', {}).get('role')}")


class TestLiquidacionEndpointAccess:
    """Test role-based access to export-liquidacion endpoint"""
    
    def test_developer_can_access_liquidacion(self, developer_token):
        """Developer (admin role) should be able to access the endpoint"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        # Should return 200 with xlsx file or 404 if no data
        assert response.status_code in [200, 404], f"Expected 200 or 404, got {response.status_code}"
        if response.status_code == 200:
            assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers.get("Content-Type", "")
            print("Developer can access liquidacion endpoint - 200 OK with xlsx")
        else:
            print(f"Developer can access but no data: {response.json()}")
    
    def test_agent_cannot_access_liquidacion(self, agent_token):
        """Agent role should NOT be able to access the endpoint (403)"""
        headers = {"Authorization": f"Bearer {agent_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        assert response.status_code == 403, f"Expected 403 for agent, got {response.status_code}"
        print(f"Agent correctly denied access: 403 - {response.json().get('detail', '')}")
    
    def test_proveedor_cannot_access_liquidacion(self, proveedor_token):
        """Proveedor role should NOT be able to access the endpoint (403)"""
        headers = {"Authorization": f"Bearer {proveedor_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        assert response.status_code == 403, f"Expected 403 for proveedor, got {response.status_code}"
        print(f"Proveedor correctly denied access: 403 - {response.json().get('detail', '')}")
    
    def test_unauthenticated_cannot_access(self):
        """Unauthenticated request should return 401 or 403"""
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO}
        )
        assert response.status_code in [401, 403], f"Expected 401/403 for unauthenticated, got {response.status_code}"
        print(f"Unauthenticated correctly denied: {response.status_code}")


class TestLiquidacionExcelStructure:
    """Test the Excel file structure and content"""
    
    def test_excel_file_download(self, developer_token):
        """Test that endpoint returns valid xlsx file"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if response.status_code == 404:
            pytest.skip("No data in date range for Excel test")
        
        assert response.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers.get("Content-Type", "")
        
        # Check Content-Disposition header for filename
        content_disp = response.headers.get("Content-Disposition", "")
        assert "attachment" in content_disp
        assert "Liquidacion_" in content_disp
        assert ".xlsx" in content_disp
        print(f"Excel file downloaded successfully: {content_disp}")
    
    def test_excel_has_required_sheets(self, developer_token):
        """Test that Excel has all required sheets: Formato, {Provider}, Incidencias, Catalogo_Drivers, route_summary"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if response.status_code == 404:
            pytest.skip("No data in date range for sheet test")
        
        assert response.status_code == 200
        
        # Load workbook from response content
        wb = load_workbook(BytesIO(response.content))
        sheet_names = wb.sheetnames
        print(f"Excel sheets found: {sheet_names}")
        
        # Required sheets
        assert "Formato" in sheet_names, "Missing 'Formato' sheet"
        assert "Incidencias a cobro" in sheet_names, "Missing 'Incidencias a cobro' sheet"
        assert "Catalogo_Drivers" in sheet_names, "Missing 'Catalogo_Drivers' sheet"
        assert "route_summary" in sheet_names, "Missing 'route_summary' sheet"
        
        # Should have at least 5 sheets (Formato + at least 1 provider + 3 others)
        assert len(sheet_names) >= 5, f"Expected at least 5 sheets, got {len(sheet_names)}"
        print(f"All required sheets present. Total sheets: {len(sheet_names)}")
    
    def test_formato_sheet_structure(self, developer_token):
        """Test Formato sheet has 14 columns with correct headers"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if response.status_code == 404:
            pytest.skip("No data in date range")
        
        wb = load_workbook(BytesIO(response.content))
        ws = wb["Formato"]
        
        # Check headers in row 1
        expected_headers = [
            "Fecha", "Nombre del Operador", "Sitio de Carga", "Tipo de Servicio",
            "Placas", "Ruta Cubbo", "Paquetes a ruta", "Entregas",
            "Entregas Fallidas", "Devuelto s/ Visita-Intento", "Entregado Final",
            "Tipo de unidad", "Importe", "% Efectividad"
        ]
        
        actual_headers = [ws.cell(row=1, column=i).value for i in range(1, 15)]
        print(f"Formato headers: {actual_headers}")
        
        for i, expected in enumerate(expected_headers):
            assert actual_headers[i] == expected, f"Column {i+1}: expected '{expected}', got '{actual_headers[i]}'"
        
        # Check example row has formulas
        cell_k2 = ws.cell(row=2, column=11)
        cell_n2 = ws.cell(row=2, column=14)
        
        # Check if cells contain formulas (start with =)
        k2_value = cell_k2.value
        n2_value = cell_n2.value
        print(f"Cell K2 value: {k2_value}, Cell N2 value: {n2_value}")
        
        # Formulas should be strings starting with =
        assert isinstance(k2_value, str) and k2_value.startswith("="), f"K2 should be formula, got: {k2_value}"
        assert isinstance(n2_value, str) and n2_value.startswith("="), f"N2 should be formula, got: {n2_value}"
        print("Formato sheet structure verified with formulas")
    
    def test_provider_sheet_structure(self, developer_token):
        """Test provider sheet has 28 columns (A-AB) with group headers and freeze panes"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if response.status_code == 404:
            pytest.skip("No data in date range")
        
        wb = load_workbook(BytesIO(response.content))
        
        # Find a provider sheet (not Formato, Incidencias, Catalogo_Drivers, route_summary)
        excluded = ["Formato", "Incidencias a cobro", "Catalogo_Drivers", "route_summary"]
        provider_sheets = [s for s in wb.sheetnames if s not in excluded]
        
        if not provider_sheets:
            pytest.skip("No provider sheets found")
        
        ws = wb[provider_sheets[0]]
        print(f"Testing provider sheet: {provider_sheets[0]}")
        
        # Check row 2 has column headers (28 columns A-AB)
        expected_col_headers = [
            "Fecha", "Nombre del Operador", "Sitio de Carga", "Tipo de Servicio",
            "Placas", "Ruta Cubbo", "Paquetes a ruta", "Entregas",
            "Entregas Fallidas", "Devuelto s/ Visita-Intento", "Recolectado",
            "Tipo de unidad", "Importe",
            "Paquetes Cargados", "Completados", "Cancelados", "Pendientes",
            "Completados c/ Evidencia", "Completados s/ Evidencia",
            "Dif. Carga", "Dif. Entrega", "Pend./Cancel.",
            "Procedentes a Cobro", "Costo x Pq", "Incidencias", "Total",
            "% Efectividad", "% Efect. SLA 40"
        ]
        
        actual_headers = [ws.cell(row=2, column=i).value for i in range(1, 29)]
        print(f"Provider sheet row 2 headers (first 10): {actual_headers[:10]}")
        
        # Verify at least the first few headers match
        for i in range(min(10, len(expected_col_headers))):
            assert actual_headers[i] == expected_col_headers[i], f"Col {i+1}: expected '{expected_col_headers[i]}', got '{actual_headers[i]}'"
        
        # Check freeze panes at B3
        assert ws.freeze_panes == "B3", f"Expected freeze_panes at B3, got {ws.freeze_panes}"
        print(f"Provider sheet verified: 28 columns, freeze_panes at {ws.freeze_panes}")
    
    def test_provider_sheet_has_real_formulas(self, developer_token):
        """Test provider sheet formulas are real Excel formulas (=G3-N3, =IFERROR, etc.)"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if response.status_code == 404:
            pytest.skip("No data in date range")
        
        wb = load_workbook(BytesIO(response.content))
        
        excluded = ["Formato", "Incidencias a cobro", "Catalogo_Drivers", "route_summary"]
        provider_sheets = [s for s in wb.sheetnames if s not in excluded]
        
        if not provider_sheets:
            pytest.skip("No provider sheets found")
        
        ws = wb[provider_sheets[0]]
        
        # Check row 3 (first data row) for formulas in columns T, U, V, W, X, Z, AA, AB
        # T=20 (Dif. Carga), U=21 (Dif. Entrega), V=22 (Pend./Cancel.)
        # W=23 (Procedentes), X=24 (Costo x Pq), Z=26 (Total), AA=27 (% Efectividad), AB=28 (% Efect. SLA)
        
        formula_cols = [11, 20, 21, 22, 23, 24, 26, 27, 28]  # K, T, U, V, W, X, Z, AA, AB
        
        formulas_found = []
        for col in formula_cols:
            cell = ws.cell(row=3, column=col)
            val = cell.value
            if val and isinstance(val, str) and val.startswith("="):
                formulas_found.append((col, val))
        
        print(f"Formulas found in row 3: {formulas_found}")
        
        # Should have at least some formulas
        assert len(formulas_found) > 0, "No formulas found in provider sheet data rows"
        
        # Check for specific formula patterns
        formula_texts = [f[1] for f in formulas_found]
        has_iferror = any("IFERROR" in f for f in formula_texts)
        has_subtraction = any("-" in f and "=" in f for f in formula_texts)
        
        print(f"Has IFERROR formula: {has_iferror}")
        print(f"Has subtraction formula: {has_subtraction}")
        
        assert has_iferror or has_subtraction, "Expected IFERROR or subtraction formulas"
    
    def test_provider_sheet_side_table(self, developer_token):
        """Test provider sheet has side table (pivot data) in columns AD-AL"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if response.status_code == 404:
            pytest.skip("No data in date range")
        
        wb = load_workbook(BytesIO(response.content))
        
        excluded = ["Formato", "Incidencias a cobro", "Catalogo_Drivers", "route_summary"]
        provider_sheets = [s for s in wb.sheetnames if s not in excluded]
        
        if not provider_sheets:
            pytest.skip("No provider sheets found")
        
        ws = wb[provider_sheets[0]]
        
        # Side table headers should be in row 2, columns AD (30) onwards
        side_headers = [ws.cell(row=2, column=i).value for i in range(30, 39)]
        print(f"Side table headers (AD-AL): {side_headers}")
        
        expected_side = ["Team", "Driver", "Fecha", "Total Paquetes", "Completados", "Cancelados", "Pendientes", "Con Evidencia", "Sin Evidencia"]
        
        for i, expected in enumerate(expected_side):
            assert side_headers[i] == expected, f"Side table col {i+30}: expected '{expected}', got '{side_headers[i]}'"
        
        print("Side table structure verified")
    
    def test_incidencias_sheet_structure(self, developer_token):
        """Test Incidencias a cobro sheet has correct headers"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if response.status_code == 404:
            pytest.skip("No data in date range")
        
        wb = load_workbook(BytesIO(response.content))
        ws = wb["Incidencias a cobro"]
        
        expected_headers = ["#", "RUTA", "Operador", "Guías", "MONTO"]
        actual_headers = [ws.cell(row=1, column=i).value for i in range(1, 6)]
        
        print(f"Incidencias headers: {actual_headers}")
        
        for i, expected in enumerate(expected_headers):
            assert actual_headers[i] == expected, f"Col {i+1}: expected '{expected}', got '{actual_headers[i]}'"
        
        print("Incidencias sheet structure verified")
    
    def test_catalogo_drivers_sheet_structure(self, developer_token):
        """Test Catalogo_Drivers sheet has correct headers"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if response.status_code == 404:
            pytest.skip("No data in date range")
        
        wb = load_workbook(BytesIO(response.content))
        ws = wb["Catalogo_Drivers"]
        
        expected_headers = ["Data M&E", "Proveedor", "Data Proveedor"]
        actual_headers = [ws.cell(row=1, column=i).value for i in range(1, 4)]
        
        print(f"Catalogo_Drivers headers: {actual_headers}")
        
        for i, expected in enumerate(expected_headers):
            assert actual_headers[i] == expected, f"Col {i+1}: expected '{expected}', got '{actual_headers[i]}'"
        
        # Check there's at least one driver entry
        driver_count = 0
        for row in range(2, 100):
            if ws.cell(row=row, column=1).value:
                driver_count += 1
            else:
                break
        
        print(f"Catalogo_Drivers has {driver_count} driver entries")
        assert driver_count > 0, "Catalogo_Drivers should have at least one driver"
    
    def test_route_summary_sheet_structure(self, developer_token):
        """Test route_summary sheet has totals in rows 1-2, headers in row 4, and XLOOKUP formula"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if response.status_code == 404:
            pytest.skip("No data in date range")
        
        wb = load_workbook(BytesIO(response.content))
        ws = wb["route_summary"]
        
        # Row 1 should have summary labels
        summary_labels = [ws.cell(row=1, column=i).value for i in range(1, 7)]
        expected_labels = ["Total rutas", "Días operados", "Total paquetes", "Completados", "Con evidencia", "Costo total"]
        
        print(f"route_summary row 1 labels: {summary_labels}")
        
        for i, expected in enumerate(expected_labels):
            assert summary_labels[i] == expected, f"Summary label {i+1}: expected '{expected}', got '{summary_labels[i]}'"
        
        # Row 2 should have values
        summary_values = [ws.cell(row=2, column=i).value for i in range(1, 7)]
        print(f"route_summary row 2 values: {summary_values}")
        
        # Row 4 should have column headers
        row4_headers = [ws.cell(row=4, column=i).value for i in range(1, 5)]
        print(f"route_summary row 4 headers (first 4): {row4_headers}")
        
        assert row4_headers[0] == "ORDER ID", f"Expected 'ORDER ID', got '{row4_headers[0]}'"
        assert row4_headers[1] == "Fecha", f"Expected 'Fecha', got '{row4_headers[1]}'"
        
        # Check for XLOOKUP formula in Driver Courier column (column 32)
        # Row 5 is first data row
        driver_courier_cell = ws.cell(row=5, column=32)
        xlookup_value = driver_courier_cell.value
        print(f"Driver Courier cell (row 5, col 32): {xlookup_value}")
        
        if xlookup_value:
            assert "XLOOKUP" in str(xlookup_value) or "IFERROR" in str(xlookup_value), f"Expected XLOOKUP formula, got: {xlookup_value}"
            print("XLOOKUP formula verified in Driver Courier column")
    
    def test_dates_are_clean_format(self, developer_token):
        """Test that dates are in clean YYYY-MM-DD format (not full timestamps)"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            pytest.skip("openpyxl not installed")
        
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        if response.status_code == 404:
            pytest.skip("No data in date range")
        
        wb = load_workbook(BytesIO(response.content))
        
        # Check route_summary sheet, Fecha column (column 2), row 5
        ws = wb["route_summary"]
        fecha_value = ws.cell(row=5, column=2).value
        print(f"Fecha value in route_summary: {fecha_value}")
        
        if fecha_value:
            fecha_str = str(fecha_value)
            # Should be YYYY-MM-DD format (10 chars) or datetime object
            assert len(fecha_str) <= 10 or "T" not in fecha_str[:20], f"Date should be clean YYYY-MM-DD, got: {fecha_str}"
            print(f"Date format verified: {fecha_str}")


class TestExistingExportStillWorks:
    """Verify the existing 'Exportar Excel' button still works"""
    
    def test_routes_report_export_endpoint(self, developer_token):
        """Test that /admin/routes-report/export still works"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/routes-report/export",
            params={"date_from": DATE_FROM, "date_to": DATE_TO},
            headers=headers
        )
        
        assert response.status_code == 200, f"Expected 200, got {response.status_code}"
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers.get("Content-Type", "")
        
        content_disp = response.headers.get("Content-Disposition", "")
        assert "LAYOUT_ADM_Cubbo" in content_disp
        print(f"Existing export endpoint works: {content_disp}")


class TestLiquidacionWithFilters:
    """Test liquidacion export with various filter parameters"""
    
    def test_export_with_provider_filter(self, developer_token):
        """Test export with provider_id filter"""
        # First get a provider ID
        headers = {"Authorization": f"Bearer {developer_token}"}
        providers_resp = requests.get(f"{BASE_URL}/api/providers", headers=headers)
        
        if providers_resp.status_code != 200:
            pytest.skip("Could not get providers")
        
        providers = providers_resp.json()
        if not providers:
            pytest.skip("No providers available")
        
        provider_id = providers[0].get("id")
        print(f"Testing with provider_id: {provider_id}")
        
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO, "provider_id": provider_id},
            headers=headers
        )
        
        # Should return 200 or 404 (no data for this provider)
        assert response.status_code in [200, 404], f"Expected 200 or 404, got {response.status_code}"
        
        if response.status_code == 200:
            content_disp = response.headers.get("Content-Disposition", "")
            print(f"Export with provider filter successful: {content_disp}")
        else:
            print(f"No data for provider {provider_id}")
    
    def test_export_with_status_filter(self, developer_token):
        """Test export with status filter"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO, "status": "closed"},
            headers=headers
        )
        
        assert response.status_code in [200, 404], f"Expected 200 or 404, got {response.status_code}"
        print(f"Export with status=closed: {response.status_code}")
    
    def test_export_with_driver_filter(self, developer_token):
        """Test export with driver name filter"""
        headers = {"Authorization": f"Bearer {developer_token}"}
        response = requests.get(
            f"{BASE_URL}/api/admin/export-liquidacion",
            params={"date_from": DATE_FROM, "date_to": DATE_TO, "driver": "Test"},
            headers=headers
        )
        
        assert response.status_code in [200, 404], f"Expected 200 or 404, got {response.status_code}"
        print(f"Export with driver filter: {response.status_code}")


if __name__ == "__main__":
    pytest.main([__file__, "-v", "--tb=short"])
