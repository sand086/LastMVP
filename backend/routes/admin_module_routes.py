"""
Admin Module — AI Consumption, Route Reporting, and Cost Configuration
New router: /api/admin/...
"""
import asyncio
import io
import logging
import uuid
from datetime import datetime, timezone, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from dependencies import db, get_current_user, require_role

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/admin", tags=["Admin Module"])

ADMIN_ROLES = ["developer", "ejecutivo", "executive", "coordinator"]
EDIT_ROLES = ["developer"]


def _require_admin(user: dict = Depends(get_current_user)):
    if user.get("role") not in ADMIN_ROLES:
        raise HTTPException(status_code=403, detail="Sin permisos para acceder al módulo admin")
    return user


def _require_editor(user: dict = Depends(get_current_user)):
    if user.get("role") not in EDIT_ROLES:
        raise HTTPException(status_code=403, detail="Solo Developer puede editar configuración")
    return user


def _month_range(period: str, date_from: str = None, date_to: str = None):
    now = datetime.now(timezone.utc)
    if period == "prev_month":
        first = (now.replace(day=1) - timedelta(days=1)).replace(day=1)
        last = now.replace(day=1) - timedelta(days=1)
        return first.strftime("%Y-%m-%dT00:00:00"), last.strftime("%Y-%m-%dT23:59:59")
    elif period == "custom" and date_from and date_to:
        return f"{date_from}T00:00:00", f"{date_to}T23:59:59"
    else:  # current_month default
        first = now.replace(day=1)
        return first.strftime("%Y-%m-%dT00:00:00"), now.strftime("%Y-%m-%dT23:59:59")


# ═══════════════ SUMMARY ═══════════════

@router.get("/summary")
async def admin_summary(
    period: str = Query("current_month"),
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    client_id: Optional[str] = None,
    user: dict = Depends(_require_admin),
):
    dt_from, dt_to = _month_range(period, date_from, date_to)
    match = {"timestamp": {"$gte": dt_from, "$lte": dt_to}}
    if client_id:
        match["client_id"] = client_id

    pipeline = [
        {"$match": match},
        {"$group": {
            "_id": "$entregable",
            "count": {"$sum": 1},
            "tokens_input": {"$sum": "$tokens_input"},
            "tokens_output": {"$sum": "$tokens_output"},
            "tokens_prompt": {"$sum": "$tokens_prompt"},
            "tokens_total": {"$sum": "$tokens_total"},
            "cost_usd": {"$sum": "$cost_usd"},
            "cost_mxn": {"$sum": "$cost_mxn"},
        }},
    ]
    results = await db.token_usage_log.aggregate(pipeline).to_list(10)

    by_entregable = {}
    grand_total_tokens = 0
    grand_cost_usd = 0
    grand_cost_mxn = 0
    grand_input = 0
    grand_output = 0
    grand_prompt = 0

    for r in results:
        ent = r["_id"]
        count = r["count"]
        avg_input = round(r["tokens_input"] / count) if count else 0
        avg_output = round(r["tokens_output"] / count) if count else 0
        avg_prompt = round(r["tokens_prompt"] / count) if count else 0

        by_entregable[ent] = {
            "count": count,
            "tokens_input": r["tokens_input"],
            "tokens_output": r["tokens_output"],
            "tokens_prompt": r["tokens_prompt"],
            "tokens_total": r["tokens_total"],
            "cost_usd": round(r["cost_usd"], 2),
            "cost_mxn": round(r["cost_mxn"], 2),
            "avg_per_unit": {"input": avg_input, "output": avg_output, "prompt": avg_prompt, "total": avg_input + avg_output + avg_prompt},
        }
        grand_total_tokens += r["tokens_total"]
        grand_cost_usd += r["cost_usd"]
        grand_cost_mxn += r["cost_mxn"]
        grand_input += r["tokens_input"]
        grand_output += r["tokens_output"]
        grand_prompt += r["tokens_prompt"]

    input_pct = round(grand_input / grand_total_tokens * 100) if grand_total_tokens else 0
    output_pct = round(grand_output / grand_total_tokens * 100) if grand_total_tokens else 0
    prompt_pct = 100 - input_pct - output_pct if grand_total_tokens else 0

    # Budget alert check
    budget_doc = await db.config.find_one({"key": "budget_alerts"}, {"_id": 0})
    budget = budget_doc.get("value", {}) if budget_doc else {}
    threshold = budget.get("monthly_threshold_usd", 50)
    over_budget = grand_cost_usd > threshold if budget.get("alert_enabled", False) else False

    return {
        "by_entregable": by_entregable,
        "totals": {
            "tokens_total": grand_total_tokens,
            "cost_usd": round(grand_cost_usd, 2),
            "cost_mxn": round(grand_cost_mxn, 2),
            "input_pct": input_pct,
            "output_pct": output_pct,
            "prompt_pct": prompt_pct,
        },
        "evaluaciones_count": by_entregable.get("evaluacion", {}).get("count", 0),
        "lumi_count": by_entregable.get("lumi", {}).get("count", 0),
        "reportes_count": by_entregable.get("reporte", {}).get("count", 0),
        "over_budget": over_budget,
        "budget_threshold": threshold,
    }


# ═══════════════ TOKEN USAGE ═══════════════

@router.get("/token-usage")
async def get_token_usage(
    period: str = Query("current_month"),
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    entregable: Optional[str] = None,
    modelo: Optional[str] = None,
    client_id: Optional[str] = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
    user: dict = Depends(_require_admin),
):
    dt_from, dt_to = _month_range(period, date_from, date_to)
    match = {"timestamp": {"$gte": dt_from, "$lte": dt_to}}
    if entregable:
        match["entregable"] = entregable
    if modelo:
        match["modelo"] = modelo
    if client_id:
        match["client_id"] = client_id

    total = await db.token_usage_log.count_documents(match)
    skip = (page - 1) * page_size

    events = await db.token_usage_log.find(
        match, {"_id": 0}
    ).sort("timestamp", -1).skip(skip).limit(page_size).to_list(page_size)

    # Summary pipeline
    summary_pipeline = [
        {"$match": match},
        {"$group": {
            "_id": "$entregable",
            "count": {"$sum": 1},
            "tokens_input": {"$sum": "$tokens_input"},
            "tokens_output": {"$sum": "$tokens_output"},
            "tokens_prompt": {"$sum": "$tokens_prompt"},
            "tokens_total": {"$sum": "$tokens_total"},
            "cost_usd": {"$sum": "$cost_usd"},
            "cost_mxn": {"$sum": "$cost_mxn"},
        }},
    ]
    summary_results = await db.token_usage_log.aggregate(summary_pipeline).to_list(10)

    by_entregable = {}
    totals = {"tokens_total": 0, "cost_usd": 0, "cost_mxn": 0, "input_pct": 0, "output_pct": 0, "prompt_pct": 0}
    ti = to = tp = 0
    for r in summary_results:
        ent = r["_id"]
        c = r["count"]
        by_entregable[ent] = {
            "count": c,
            "tokens_input": r["tokens_input"],
            "tokens_output": r["tokens_output"],
            "tokens_prompt": r["tokens_prompt"],
            "tokens_total": r["tokens_total"],
            "cost_usd": round(r["cost_usd"], 4),
            "cost_mxn": round(r["cost_mxn"], 4),
            "avg_per_unit": {
                "input": round(r["tokens_input"] / c) if c else 0,
                "output": round(r["tokens_output"] / c) if c else 0,
                "prompt": round(r["tokens_prompt"] / c) if c else 0,
                "total": round(r["tokens_total"] / c) if c else 0,
            },
        }
        totals["tokens_total"] += r["tokens_total"]
        totals["cost_usd"] += r["cost_usd"]
        totals["cost_mxn"] += r["cost_mxn"]
        ti += r["tokens_input"]
        to += r["tokens_output"]
        tp += r["tokens_prompt"]

    tt = ti + to + tp
    totals["cost_usd"] = round(totals["cost_usd"], 4)
    totals["cost_mxn"] = round(totals["cost_mxn"], 4)
    totals["input_pct"] = round(ti / tt * 100) if tt else 0
    totals["output_pct"] = round(to / tt * 100) if tt else 0
    totals["prompt_pct"] = 100 - totals["input_pct"] - totals["output_pct"] if tt else 0

    return {
        "summary": {"by_entregable": by_entregable, "totals": totals},
        "events": events,
        "pagination": {"total": total, "page": page, "page_size": page_size, "total_pages": max(1, (total + page_size - 1) // page_size)},
    }


def _calc_working_hours(arrival: str, departure: str) -> float:
    """Calculate hours worked from arrival/departure time strings."""
    if not arrival or not departure:
        return 0
    try:
        t1 = datetime.strptime(arrival, "%H:%M")
        t2 = datetime.strptime(departure, "%H:%M")
        return round((t2 - t1).total_seconds() / 3600, 2)
    except Exception:
        return 0


def _check_on_time(arrival: str, scheduled: str) -> str:
    """Check if arrival was within 20min tolerance of scheduled time."""
    if not arrival or not scheduled:
        return ""
    try:
        t_arr = datetime.strptime(arrival, "%H:%M")
        t_sched = datetime.strptime(scheduled, "%H:%M")
        tolerance = t_sched + timedelta(minutes=20)
        return "Si" if t_arr <= tolerance else "No"
    except Exception:
        return ""


async def _build_report_row(j: dict, prov: dict) -> dict:
    """Build a single report row from a journey and its provider."""
    close_data = j.get("close_data") or {}
    start_data = j.get("start_data") or {}
    km = close_data.get("km_traveled", 0)
    km_excedente = max(0, km - 120) if km else 0

    # Quality score from packages
    pkgs = await db.packages.find(
        {"journey_id": j["id"], "evidence_score": {"$ne": None}},
        {"_id": 0, "evidence_score": 1}
    ).to_list(1000)
    scores = [p["evidence_score"] for p in pkgs]
    score_ia = round(sum(scores) / len(scores)) if scores else None
    delivered = j.get("packages_delivered", 0)
    with_evidence = sum(1 for s in scores if s == 100)

    arrival = start_data.get("arrival_time", "")
    departure = close_data.get("departure_time", "")
    scheduled = start_data.get("scheduled_time", "07:00")
    prov_team = prov.get("team", "")

    return {
        "order_id": j.get("cosmo_route_id", j.get("id", "")[:14]),
        "fecha": j.get("date", ""),
        "driver": j.get("driver_name", ""),
        "team": prov_team,
        "tipo_unidad": prov.get("vehicle_type", "Sedán"),
        "estado": prov.get("state", "CDMX/EDOMEX"),
        "tipo_servicio": prov.get("service_type", "Última milla"),
        "proveedor": prov.get("name", ""),
        "costo": prov.get("cost", 0),
        "pv": prov.get("sale_price", 0),
        "horario_asistencia": scheduled,
        "hora_entrada": arrival,
        "hora_salida": departure,
        "tolerancia": "",
        "asistencia_en_tiempo": _check_on_time(arrival, scheduled),
        "horas_laboradas": _calc_working_hours(arrival, departure),
        "distancia_km": round(km, 2),
        "km_excedente": round(km_excedente, 2) if km_excedente > 0 else None,
        "tipo_tarifa": "Tarifa extra" if km_excedente > 0 else "Tarifa normal",
        "costo_km_adicional": round(km_excedente * 5, 2) if km_excedente > 0 else None,
        "backup_activado": j.get("backup_activated", False),
        "hora_inicio_backup": j.get("backup_start_time"),
        "horas_laboradas_backup": None,
        "total_paquetes": j.get("packages_total", 0),
        "completados": delivered,
        "cancelados": j.get("packages_cancelled", 0),
        "pendientes": j.get("packages_pending", j.get("packages_total", 0) - delivered - j.get("packages_failed", 0)),
        "con_evidencia": with_evidence,
        "sin_evidencia": max(0, delivered - with_evidence),
        "score_ia": score_ia,
        "comentarios": j.get("comments", ""),
        "journey_id": j["id"],
    }


# ═══════════════ ROUTES REPORT ═══════════════

@router.get("/routes-report")
async def routes_report(
    date_from: str = Query(...),
    date_to: str = Query(...),
    driver: Optional[str] = None,
    team: Optional[str] = None,
    provider_id: Optional[str] = None,
    status: Optional[str] = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
    user: dict = Depends(_require_admin),
):
    j_query = {"date": {"$gte": date_from, "$lte": date_to}}
    if provider_id:
        j_query["provider_id"] = provider_id
    if status:
        j_query["status"] = status
    if driver:
        j_query["driver_name"] = {"$regex": driver, "$options": "i"}

    journeys = await db.journeys.find(j_query, {"_id": 0}).sort("date", -1).to_list(5000)

    providers_map = {}
    for p in await db.providers.find({}, {"_id": 0}).to_list(100):
        providers_map[p["id"]] = p

    rows = []
    for j in journeys:
        prov = providers_map.get(j.get("provider_id"), {})
        prov_team = prov.get("team", "")
        if team and team.lower() not in prov_team.lower():
            continue
        row = await _build_report_row(j, prov)
        rows.append(row)

    total = len(rows)
    skip = (page - 1) * page_size
    paged_rows = rows[skip:skip + page_size]

    # Totals
    totals = {
        "total_rutas": total,
        "total_dias": len(set(r["fecha"] for r in rows if r["fecha"])),
        "total_paquetes": sum(r["total_paquetes"] for r in rows),
        "completados": sum(r["completados"] for r in rows),
        "con_evidencia": sum(r["con_evidencia"] for r in rows),
        "costo_total": sum(r["costo"] or 0 for r in rows),
        "km_excedente_cost": sum(r["costo_km_adicional"] or 0 for r in rows),
    }

    return {
        "totals": totals,
        "rows": paged_rows,
        "pagination": {"total": total, "page": page, "page_size": page_size, "total_pages": max(1, (total + page_size - 1) // page_size)},
    }


# ═══════════════ ROUTES REPORT EXCEL EXPORT ═══════════════

@router.get("/routes-report/export")
async def export_routes_report(
    date_from: str = Query(...),
    date_to: str = Query(...),
    driver: Optional[str] = None,
    team: Optional[str] = None,
    provider_id: Optional[str] = None,
    status: Optional[str] = None,
    columns: str = "",
    user: dict = Depends(_require_admin),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    # Reuse the routes-report logic
    data = await routes_report(
        date_from=date_from, date_to=date_to, driver=driver, team=team,
        provider_id=provider_id, status=status, page=1, page_size=99999, user=user,
    )

    all_cols = [
        ("order_id", "ORDER ID"), ("fecha", "Fecha"), ("driver", "Driver"), ("team", "Team"),
        ("tipo_unidad", "Tipo de unidad"), ("estado", "Estado"), ("tipo_servicio", "Tipo de servicio"),
        ("proveedor", "Proveedor"), ("costo", "Costo"), ("pv", "PV"),
        ("horario_asistencia", "Horario de asistencia"), ("hora_entrada", "Hora de entrada"),
        ("hora_salida", "Hora de salida"), ("asistencia_en_tiempo", "Asistencia en tiempo"),
        ("horas_laboradas", "Horas laboradas"), ("distancia_km", "Distancia (KM)"),
        ("km_excedente", "KM excedente"), ("tipo_tarifa", "Tipo de tarifa"),
        ("costo_km_adicional", "Costo por KM adicional"), ("backup_activado", "¿Se activó backup?"),
        ("hora_inicio_backup", "Hora inicio backup"), ("horas_laboradas_backup", "Horas laboradas backup"),
        ("total_paquetes", "Total de paquetes"), ("completados", "Completados"),
        ("cancelados", "Cancelados"), ("pendientes", "Pendientes"),
        ("con_evidencia", "Completados con evidencia"), ("sin_evidencia", "Completados sin evidencia"),
        ("score_ia", "Score IA"), ("comentarios", "Comentarios"),
    ]

    selected = columns.split(",") if columns else [c[0] for c in all_cols]
    cols = [(key, label) for key, label in all_cols if key in selected]

    wb = Workbook()
    ws = wb.active
    ws.title = "route_summary"

    header_fill = PatternFill(start_color="1A1916", end_color="1A1916", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    for col_idx, (key, label) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row_idx, row in enumerate(data["rows"], 2):
        for col_idx, (key, _) in enumerate(cols, 1):
            val = row.get(key)
            if val is None:
                val = ""
            elif isinstance(val, bool):
                val = "Si" if val else "No"
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-width
    for col in ws.columns:
        max_len = 0
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 30)

    # Totals sheet
    ws2 = wb.create_sheet("totales")
    totals = data["totals"]
    totals_headers = ["Total rutas", "Días operados", "Total paquetes", "Completados", "Con evidencia", "Costo total", "KM excedente cost"]
    totals_values = [totals["total_rutas"], totals["total_dias"], totals["total_paquetes"], totals["completados"], totals["con_evidencia"], totals["costo_total"], totals["km_excedente_cost"]]
    for i, (h, v) in enumerate(zip(totals_headers, totals_values), 1):
        ws2.cell(row=1, column=i, value=h).font = Font(bold=True)
        ws2.cell(row=2, column=i, value=v)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"LAYOUT_ADM_Cubbo_{date_from}_{date_to}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ═══════════════ CONFIG ═══════════════

@router.get("/config")
async def get_admin_config(user: dict = Depends(_require_admin)):
    ia_cost = await db.config.find_one({"key": "ia_cost_config"}, {"_id": 0})
    exchange = await db.config.find_one({"key": "exchange_rate"}, {"_id": 0})
    budget = await db.config.find_one({"key": "budget_alerts"}, {"_id": 0})

    return {
        "ia_cost_config": ia_cost.get("value", {}) if ia_cost else {"models": []},
        "exchange_rate": exchange.get("value", {"rate": 19.0, "source": "manual", "auto_update": False, "history": []}) if exchange else {"rate": 19.0, "source": "manual", "auto_update": False, "history": []},
        "budget_alerts": budget.get("value", {"monthly_threshold_usd": 50, "alert_enabled": False, "weekly_report_enabled": False}) if budget else {"monthly_threshold_usd": 50, "alert_enabled": False, "weekly_report_enabled": False},
    }


@router.patch("/config")
async def update_admin_config(
    payload: dict,
    user: dict = Depends(_require_editor),
):
    section = payload.get("section")
    value = payload.get("value")

    if section not in ("ia_cost_config", "exchange_rate", "budget_alerts"):
        raise HTTPException(status_code=400, detail="Sección inválida")

    now = datetime.now(timezone.utc).isoformat()

    # If updating exchange_rate, append to history
    if section == "exchange_rate" and value.get("rate"):
        existing = await db.config.find_one({"key": "exchange_rate"}, {"_id": 0})
        history = existing.get("value", {}).get("history", []) if existing else []
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        history = [h for h in history if h.get("date") != today]
        history.insert(0, {"date": today, "rate": value["rate"]})
        value["history"] = history[:30]
        value["last_updated"] = now

    await db.config.update_one(
        {"key": section},
        {"$set": {"value": value, "updated_by": user.get("name", user["email"]), "updated_at": now}},
        upsert=True,
    )

    # Audit log
    await db.audit_logs.insert_one({
        "id": str(uuid.uuid4()),
        "action": f"admin_config_update:{section}",
        "user": user.get("name", user["email"]),
        "role": user.get("role"),
        "timestamp": now,
        "details": {"section": section},
    })

    return {"success": True, "section": section}
